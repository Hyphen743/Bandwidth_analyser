from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

OUT = Path('Bandwidth_Analyzer_Review_1_Complete.pptx')
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]

BG = 'F8F5EF'; INK = '292824'; MUTED = '706D66'; TERRA = 'C8553D'
AMBER = 'D99A42'; SAGE = '718970'; PALE = 'E8E1D7'; WHITE = 'FFFFFF'; DARK = '252421'
FONT = 'Aptos'

def rgb(v): return RGBColor.from_string(v)

def box(slide, x,y,w,h, color, shape=MSO_SHAPE.RECTANGLE, line=None, radius=False):
    typ = MSO_SHAPE.ROUNDED_RECTANGLE if radius else shape
    sh=slide.shapes.add_shape(typ, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=rgb(color)
    if line:
        sh.line.color.rgb=rgb(line); sh.line.width=Pt(1)
    else: sh.line.fill.background()
    if radius:
        try: sh.adjustments[0]=0.10
        except Exception: pass
    return sh

def txt(slide,x,y,w,h,content,size=18,color=INK,bold=False,align=PP_ALIGN.LEFT,
        valign=MSO_ANCHOR.MIDDLE,margin=0.02,font=FONT,italic=False):
    sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=sh.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(margin)
    tf.margin_top=tf.margin_bottom=Inches(margin)
    tf.vertical_anchor=valign
    for j,line in enumerate(content.split('\n')):
        p=tf.paragraphs[0] if j==0 else tf.add_paragraph()
        p.alignment=align; p.space_after=Pt(0); p.space_before=Pt(0)
        p.line_spacing=1.04
        r=p.add_run(); r.text=line
        r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.italic=italic
        r.font.color.rgb=rgb(color)
    return sh

def ln(slide,x1,y1,x2,y2,color=PALE,width=1.2):
    sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb=rgb(color); sh.line.width=Pt(width); return sh

def dot(slide,cx,cy,d,color):
    sh=slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx-d/2), Inches(cy-d/2), Inches(d), Inches(d))
    sh.fill.solid(); sh.fill.fore_color.rgb=rgb(color); sh.line.fill.background(); return sh

def base(section,title,subtitle=None,dark=False):
    s=prs.slides.add_slide(blank); bg=DARK if dark else BG; fg=BG if dark else INK
    muted='BDB7AE' if dark else MUTED
    s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb(bg)
    box(s,0,0,.12,7.5,TERRA)
    txt(s,.58,.25,10,.26,section.upper(),10,TERRA,True)
    txt(s,.58,.63,12.1,.68,title,29,fg,True)
    ln(s,.58,1.47,12.76,1.47,'4C4943' if dark else PALE,1)
    if subtitle: txt(s,.60,1.55,12,.42,subtitle,13,muted)
    txt(s,.58,7.12,7,.2,'BAMAT207  ·  PROBABILITY & STATISTICS',9,muted,True)
    txt(s,11.95,7.12,.8,.2,f'{len(prs.slides):02d} / 10',9,muted,True,PP_ALIGN.RIGHT)
    return s

# 1: Cover
s=prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb(BG)
box(s,0,0,.14,7.5,TERRA); box(s,8.40,0,4.93,7.5,DARK)
txt(s,.72,.55,6.8,.35,'BAMAT207  /  REVIEW 01',12,TERRA,True)
txt(s,.72,1.18,7.05,1.50,'Reading the\nshape of network traffic',35,INK,True)
txt(s,.75,2.95,6.90,.77,'A statistical bandwidth analyzer for seeing demand,\nspotting unusual activity and planning capacity.',17,MUTED)
ln(s,.75,4.02,7.8,4.02,PALE,1)
txt(s,.75,4.25,3,.24,'PROJECT BY',9,MUTED,True)
txt(s,.75,4.53,6.7,.43,'Raman Darda  ·  25BCE0893',20,INK,True)
txt(s,.75,5.30,3,.24,'FACULTY GUIDE',9,MUTED,True)
txt(s,.75,5.59,6.7,.43,'Prof. Juhi Baroi',18,INK,True)
# Editable abstract signal artwork
for d,c in [(2.05,'383732'),(1.48,'4C4942'),(.90,TERRA)]:
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10.82-d/2), Inches(1.62-d/2), Inches(d), Inches(d))
    sh.fill.background(); sh.line.color.rgb=rgb(c); sh.line.width=Pt(2)
txt(s,9.15,2.83,3.3,.65,'FROM RAW SAMPLES\nTO A CLEARER SIGNAL',16,BG,True,PP_ALIGN.CENTER)
for i,(h,c) in enumerate([(0.54,TERRA),(1.06,AMBER),(.74,SAGE),(1.34,TERRA),(.86,AMBER),(1.16,SAGE),(.62,TERRA),(1.0,AMBER)]):
    box(s,8.87+i*.46,5.20-h,.20,h,c,radius=True)
txt(s,9.0,5.62,3.8,.5,'Idea  ·  problem  ·  method  ·  plan',11,'BDB7AE',False,PP_ALIGN.CENTER)

# 2: Problem and study questions
s=base('01 / The problem','A connection can feel slow before its cause is clear.','The same symptom can come from different traffic patterns.')
txt(s,.72,2.15,4.6,1.15,'“Why is the network\nslow right now?”',25,INK,True)
txt(s,.72,3.60,4.55,1.40,'A single usage total cannot tell us who used bandwidth, when demand peaked, or whether latency and packet loss rose at the same time.',16,MUTED)
ln(s,5.83,2.05,5.83,6.10,PALE,1)
for i,(n,h,b) in enumerate([('01','WHO','Which device or app is using the link?'),('02','WHEN','Are peaks tied to time of day or day of week?'),('03','WHAT CHANGED','Is a spike routine demand or an anomaly?'),('04','WHAT NEXT','Will demand approach available capacity?')]):
    y=2.02+i*.98; txt(s,6.22,y,.54,.30,n,11,TERRA,True); txt(s,6.94,y,4.7,.30,h,13,INK,True); txt(s,6.94,y+.32,5.45,.43,b,14,MUTED)

# 3: Project concept and objectives
s=base('02 / Project idea','Turn measurements into network decisions.','Four questions organize the statistical study.')
items=[('DESCRIBE','How does demand vary by hour, day, application and device?'),('DETECT','Which usage intervals depart from their expected pattern?'),('FORECAST','How closely can recent history predict future demand?'),('PLAN','When might demand put pressure on available capacity?')]
for i,(h,b) in enumerate(items):
    col=i%2; row=i//2; x=.78+col*6.08; y=2.15+row*1.72
    box(s,x,y,.06,1.10,[TERRA,AMBER,SAGE,TERRA][i]); txt(s,x+.25,y,5,.32,h,12,[TERRA,AMBER,SAGE,TERRA][i],True)
    txt(s,x+.25,y+.40,5.25,.65,b,16,INK,True)
box(s,.78,5.83,11.75,.55,'EEE8DE',radius=True)
txt(s,1.02,5.91,11.2,.34,'Aim: build an interpretable path from hourly measurements to evidence and capacity planning.',13,INK,True)

# 4: Architecture flowchart + stack
s=base('03 / System design','A simple pipeline connects collection to insight.','Each stage uses tools already present in the repository.')
flow=[('COLLECT','psutil · ping\nscapy optional',TERRA),('STORE','SQLite\nstructured tables',AMBER),('PREPARE','Python · Pandas\nNumPy',SAGE),('MODEL','Scikit-learn\nStatsmodels',TERRA),('PRESENT','Streamlit\nPlotly',AMBER)]
for i,(h,b,c) in enumerate(flow):
    x=.60+i*2.56
    box(s,x,2.40,2.12,1.28,'FFFFFF',line=PALE,radius=True)
    box(s,x,2.40,2.12,.07,c)
    txt(s,x+.13,2.62,1.85,.30,h,12,c,True,PP_ALIGN.CENTER)
    txt(s,x+.10,3.00,1.92,.55,b,13,INK,True,PP_ALIGN.CENTER)
    if i<4:
        ln(s,x+2.15,3.04,x+2.47,3.04,MUTED,1.3)
        txt(s,x+2.18,2.78,.28,.36,'›',18,TERRA,True,PP_ALIGN.CENTER)
txt(s,.78,4.12,2.0,.30,'INPUT DATA',11,TERRA,True)
txt(s,2.30,4.08,9.85,.52,'Simulated 90-day traffic (5 users, 8 devices, 12 apps) or local live counters; bytes, timestamps, latency, jitter and loss.',13,MUTED)
ln(s,.78,4.76,12.50,4.76,PALE,1)
txt(s,.78,5.02,2.35,.30,'REPORTING & ALERTS',11,TERRA,True)
txt(s,3.00,4.97,9.50,.58,'ReportLab exports PDF reports · Plyer can send desktop alerts · SQLite stores anomaly labels and forecast records.',13,MUTED)
box(s,.78,5.80,11.75,.58,DARK,radius=True)
txt(s,1.02,5.89,11.15,.34,'STACK AT A GLANCE   Python · Streamlit · SQLite · Pandas · NumPy · scikit-learn · Statsmodels · Plotly · psutil',12,BG,True)

# 5: Statistical anomaly formulas
s=base('04 / Anomalies — statistical rules','Z-score and IQR compare traffic with a seasonal baseline.','The code groups by hour of day and weekday/weekend, then flags unusually high traffic.')
txt(s,.82,2.07,2.15,.32,'Z-SCORE',12,TERRA,True)
box(s,.80,2.49,5.65,.69,'EEE8DE',radius=True)
txt(s,1.06,2.60,5.10,.42,'z = (x − μ) / σ',22,INK,True,PP_ALIGN.CENTER)
txt(s,.84,3.32,5.52,.64,'x = current Mbps; μ and σ = mean and standard deviation for its seasonal group.',13,MUTED)
txt(s,.84,4.08,5.52,.68,'Flag when z ≥ 2.5. Severity: info ≥2.5, warning ≥3.0, critical ≥4.5.',13,INK,True)
ln(s,6.66,2.05,6.66,5.90,PALE,1)
txt(s,6.98,2.07,2.15,.32,'IQR FENCE',12,AMBER,True)
box(s,6.95,2.49,5.55,.69,'EEE8DE',radius=True)
txt(s,7.20,2.60,5.05,.42,'IQR = Q₃ − Q₁\nUpper = Q₃ + 1.5 × IQR',18,INK,True,PP_ALIGN.CENTER)
txt(s,6.99,3.55,5.44,.70,'Q₁ and Q₃ are the 25th and 75th percentiles for that hour/weekend group.',13,MUTED)
txt(s,6.99,4.38,5.44,.76,'Flag when x > Upper. This implementation detects high spikes; it does not flag unusually low traffic.',13,INK,True)
box(s,.80,5.55,11.70,.61,DARK,radius=True)
txt(s,1.04,5.66,11.20,.34,'Both are explainable statistical rules. Their thresholds control sensitivity; they do not learn from anomaly labels.',13,BG)

# 6: Isolation forest and detector evaluation
s=base('05 / Anomalies — Isolation Forest','Isolation Forest uses random splits to isolate unusual points.','It is an unsupervised machine-learning method built from a mathematical isolation score.')
txt(s,.82,2.06,5.30,.34,'FEATURES USED BY THIS PROJECT',12,SAGE,True)
txt(s,.82,2.51,5.40,.75,'Mbps · bytes up · bytes down\nhour of day · day of week',17,INK,True)
txt(s,.82,3.54,5.30,.86,'Shorter average path to isolate a point implies it is more anomalous.',14,MUTED)
box(s,.82,4.53,5.46,.91,'EEE8DE',radius=True)
txt(s,1.03,4.59,5.02,.36,'s(x,n) = 2^(−E[h(x)] / c(n))',16,INK,True,PP_ALIGN.CENTER)
txt(s,1.03,4.97,5.02,.36,'c(n) = 2H(n−1) − 2(n−1)/n',14,INK,True,PP_ALIGN.CENTER)
txt(s,.84,5.50,5.43,.45,'Contamination 3% sets expected flagged share; severity uses score cutoffs.',11,MUTED)
ln(s,6.60,2.05,6.60,6.00,PALE,1)
txt(s,6.95,2.06,4.75,.34,'VALIDATION AGAINST SYNTHETIC LABELS',12,TERRA,True)
txt(s,6.95,2.53,5.42,.34,'Precision = TP / (TP + FP)',15,INK,True)
txt(s,6.95,3.05,5.42,.34,'Recall = TP / (TP + FN)',15,INK,True)
txt(s,6.95,3.57,5.42,.34,'F1 = 2 × Precision × Recall / (Precision + Recall)',14,INK,True)
txt(s,6.95,4.24,5.42,.82,'TP: correctly flagged · FP: false alarm · FN: missed labeled event. Higher precision means fewer false alarms; higher recall means fewer misses.',12,MUTED)
box(s,6.95,5.30,5.46,.68,'EEE8DE',radius=True)
txt(s,7.18,5.39,5.00,.48,'Isolation score in code = −Scikit-learn decision_function(X)',11,INK,True,PP_ALIGN.CENTER)

# 7: Forecast data + baseline + Holt-Winters
s=base('06 / Forecast — time-series methods','Forecasting starts with an hourly Mbps series.','Traffic is grouped by hour, then compared using a daily baseline and additive Holt–Winters.')
box(s,.80,2.00,11.70,.62,'EEE8DE',radius=True)
txt(s,1.02,2.10,11.22,.40,'Mbps = (bytes_up + bytes_down) × 8 / (3,600 × 1,000,000)',17,INK,True,PP_ALIGN.CENTER)
txt(s,.83,2.91,5.30,.31,'SEASONAL NAÏVE BASELINE',12,TERRA,True)
txt(s,.83,3.35,5.30,.46,'ŷₜ = yₜ₋₂₄',21,INK,True)
txt(s,.83,3.87,5.30,.53,'Predict this hour using the observed value from 24 hours earlier.',13,MUTED)
ln(s,6.60,2.86,6.60,5.82,PALE,1)
txt(s,6.95,2.91,5.30,.31,'ADDITIVE HOLT–WINTERS · SEASON LENGTH m = 24',12,AMBER,True)
txt(s,6.98,3.34,5.42,.38,'Level:  ℓₜ = α(yₜ − sₜ₋₂₄) + (1−α)(ℓₜ₋₁ + bₜ₋₁)',12,INK,True)
txt(s,6.98,3.86,5.42,.38,'Trend:  bₜ = β(ℓₜ − ℓₜ₋₁) + (1−β)bₜ₋₁',12,INK,True)
txt(s,6.98,4.38,5.42,.38,'Season: sₜ = γ(yₜ − ℓₜ) + (1−γ)sₜ₋₂₄',12,INK,True)
txt(s,6.98,4.91,5.42,.43,'Forecast: ŷₜ₊ₕ = ℓₜ + h·bₜ + matching seasonal effect',12,TERRA,True)
txt(s,6.98,5.43,5.40,.38,'α, β, γ control update speed for level, trend and season.',11,MUTED)
box(s,.80,5.90,11.70,.43,DARK,radius=True)
txt(s,1.02,5.94,11.22,.32,'“Additive” means level + trend + seasonal effect. This is time-series smoothing, not a straight-line regression equation.',12,BG,True)

# 8: Regression forecast + interval/capacity
s=base('07 / Forecast — regression model','Gradient Boosting learns patterns from lag and calendar features.','The dashboard uses this model to produce the next 24-hour forecast.')
txt(s,.82,2.02,5.35,.31,'FEATURE VECTOR xₜ',12,SAGE,True)
txt(s,.82,2.44,5.48,.95,'lag₁, lag₂, lag₂₄ · rolling mean/std (24h)\nhour of day · weekday · weekend flag',15,INK,True)
box(s,.82,3.59,5.48,.77,'EEE8DE',radius=True)
txt(s,1.05,3.71,5.02,.49,'Fₘ(x) = Fₘ₋₁(x) + ηhₘ(x)',18,INK,True,PP_ALIGN.CENTER)
txt(s,.84,4.51,5.42,.75,'Each added decision tree hₘ corrects patterns left in earlier predictions. The model is HistGradientBoostingRegressor.',13,MUTED)
ln(s,6.60,2.00,6.60,5.75,PALE,1)
txt(s,6.95,2.02,5.30,.31,'FORECAST BAND & CAPACITY RULE',12,TERRA,True)
txt(s,6.98,2.47,5.38,.46,'Lower / upper = ŷ ± 1.96σ₇₂',16,INK,True)
txt(s,6.98,3.01,5.38,.77,'σ₇₂ is the standard deviation of recent 72-hour traffic; lower bound is clipped at zero.',13,MUTED)
txt(s,6.98,3.99,5.38,.46,'Capacity warning if ŷ ≥ 0.85 × link capacity',14,INK,True)
txt(s,6.98,4.52,5.38,.77,'The band is a recent-volatility estimate, not a calibrated 95% model confidence interval.',13,TERRA,True)
box(s,.82,5.57,11.70,.58,'EEE8DE',radius=True)
txt(s,1.05,5.66,11.10,.36,'Validation uses the final 168 hours (7 days) as a time-ordered holdout; earlier hours train the models.',13,INK,True)

# 9: Forecast evaluation formulas / R² target
s=base('08 / Forecast — evaluation','Use error metrics and R² to describe forecast quality.','Lower error is better; R² describes fit relative to predicting the test-set mean.')
txt(s,.82,2.02,5.48,.34,'ERROR METRICS · y = actual, ŷ = predicted',12,TERRA,True)
txt(s,.85,2.50,5.40,.42,'MAE  = (1/n) Σ |yᵢ − ŷᵢ|',16,INK,True)
txt(s,.85,3.13,5.40,.42,'RMSE = √[(1/n) Σ (yᵢ − ŷᵢ)²]',16,INK,True)
txt(s,.85,3.76,5.40,.52,'MAPE = (100/n) Σ |(yᵢ − ŷᵢ) / yᵢ|',15,INK,True)
txt(s,.85,4.45,5.40,.70,'MAE: average miss in Mbps · RMSE penalizes large misses · MAPE expresses miss as a percent (undefined at actual zero).',12,MUTED)
ln(s,6.60,2.00,6.60,5.78,PALE,1)
txt(s,6.95,2.02,5.30,.34,'COEFFICIENT OF DETERMINATION',12,SAGE,True)
box(s,6.95,2.50,5.45,.72,'EEE8DE',radius=True)
txt(s,7.13,2.62,5.08,.46,'R² = 1 − Σ(yᵢ−ŷᵢ)² / Σ(yᵢ−ȳ)²',15,INK,True,PP_ALIGN.CENTER)
txt(s,6.98,3.43,5.40,.62,'R² = 1 means a perfect match on this test set; R² = 0 matches the mean-only baseline.',12,MUTED)
txt(s,6.98,4.20,5.40,.45,'PROJECT TARGET: R² > 0.85',15,TERRA,True)
txt(s,6.98,4.69,5.40,.65,'To be added and measured later. It is not currently calculated by the forecast page and is not a result yet.',12,INK,True)
box(s,.82,5.63,11.58,.53,DARK,radius=True)
txt(s,1.04,5.72,11.12,.32,'Important: high R² alone does not guarantee small forecast errors or prove the model will generalize.',12,BG,True)

# 10: Roadmap
s=base('09 / Review plan','From proposal to tested results in three reviews.','This first review establishes the mathematical approach and implementation map.')
road=[('REVIEW 01','Idea · formulas · architecture','Define the problem, methods, tech stack and evaluation plan.',TERRA),('REVIEW 02','Implementation · evidence','Show working screens, calculate metrics and tune methods.',AMBER),('REVIEW 03','Validation · conclusion','Compare results, explain limits and present findings.',SAGE)]
for i,(h,b,d,c) in enumerate(road):
    x=.79+i*4.20; box(s,x,2.46,3.22,.06,c); txt(s,x,2.74,3.35,.35,h,12,c,True)
    txt(s,x,3.23,3.50,.60,b,17,INK,True); txt(s,x,4.02,3.45,1.02,d,14,MUTED)
box(s,.78,5.42,11.75,.77,'EEE8DE',radius=True)
txt(s,1.02,5.53,11.15,.53,'DATA NOTE   Simulator = labeled, repeatable study data. Live capture = local measurements; short live history cannot support the 7-day holdout evaluation yet.',12,INK,True)
box(s,.78,6.32,11.75,.40,DARK,radius=True)
txt(s,1.02,6.35,11.15,.30,'Measure  →  explain  →  anticipate  →  plan',13,BG,True,PP_ALIGN.CENTER)

prs.save(OUT)
print(f'Created {OUT.resolve()} · {len(prs.slides)} slides · {OUT.stat().st_size:,} bytes')
